import streamlit as st
import openpyxl
from math import comb

EXCEL_PATH = "PE_data.xlsx"  # put app.py and PE_data.xlsx in same folder

@st.cache_data
def load_qty_ae_lookup(path: str) -> dict[int, int]:
    """
    Reads Sheet1 col A (customer accts) -> col B (Qty AE) mapping from the workbook.
    This is your 'automated hardcoded value' / lookup table.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]

    lookup = {}
    # Data starts at row 3 in your file and runs through 50 customer accounts
    for r in range(3, 200):
        a = ws.cell(row=r, column=1).value  # Column A
        b = ws.cell(row=r, column=2).value  # Column B
        if isinstance(a, (int, float)) and isinstance(b, (int, float)):
            lookup[int(a)] = int(b)

    if not lookup:
        raise ValueError("Could not find A->B lookup values in the Excel file.")

    return lookup

def safe_div(a: float, b: float) -> float:
    return a / b if b else 0.0

st.set_page_config(page_title="SW CO Network Effect", layout="wide")
st.title("SW CO Network effect")

lookup = load_qty_ae_lookup(EXCEL_PATH)

# ----- Inputs (your “green shaded columns”) -----
st.sidebar.header("Inputs ")

cust_accts = st.sidebar.selectbox(
    "Qty of Customer Accts ",
    options=sorted(lookup.keys()),
    index=0
)

qty_ae_per_pack = st.sidebar.selectbox(
    "Qty AE per pack ",
    options=list(range(6, 13)),
    index=0
)

# ----- Derived / hardcoded -----
qty_ae = lookup[cust_accts]  # Column B (hardcoded/lookup)

cust_acct_per_ae = safe_div(cust_accts, qty_ae)  # Column C = A/B
hunter_accts_per_ae = 2 * cust_acct_per_ae       # Column D = 2*C

links_per_pack = comb(qty_ae_per_pack, 2)        # Column F = COMBIN(E,2)
ae_links_in_each_pack = qty_ae_per_pack - 1      # Column G = E-1

ae_total_links_across_packs = (cust_acct_per_ae + hunter_accts_per_ae ) * ae_links_in_each_pack # Column H = (C+D)*G
company_wide_links = cust_accts * links_per_pack                 # Column I = A*F

# ----- Display -----
st.subheader("Results")

col1, col2, col3, col4 = st.columns(4)
col1.metric("Qty of Customer Accts ", cust_accts)
col2.metric("Qty AE ", qty_ae)
col3.metric("cust acct per AE ", f"{cust_acct_per_ae:.2f}")
col4.metric("qty hunter accts per AE ", f"{hunter_accts_per_ae:.2f}")

col5, col6, col7, col8 = st.columns(4)
col5.metric("Qty AE per pack ", qty_ae_per_pack)
col6.metric("Links per pack ", links_per_pack)
col7.metric("AE links in each pack ", ae_links_in_each_pack)
col8.metric("AE total links across packs ", f"{ae_total_links_across_packs:.2f}")

st.metric("Company-wide network connections ", company_wide_links)

st.divider()
st.subheader("Single-row table")

st.dataframe([{
    "Qty of Customer Accts ": cust_accts,
    "Qty AE (B)": qty_ae,
    "cust acct per AE ": round(cust_acct_per_ae, 2),
    "qty hunter accts per AE ": round(hunter_accts_per_ae, 2),
    "Qty AE per pack ": qty_ae_per_pack,
    "Links per pack )": links_per_pack,
    "AE links in each pack ": ae_links_in_each_pack,
    "AE total links across packs ": round(ae_total_links_across_packs, 2),
    "Company-wide links ": company_wide_links,
}])
